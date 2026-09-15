import os
import sys
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

DEMO_DIR = r"C:\Users\tonda\Desktop\demoimages"
MD_PATH = r"c:\Users\tonda\Desktop\hrp\Presentations\Reports\CSE499A_Final_Report_CardioRisk_AI.md"
DOCX_DESKTOP = r"C:\Users\tonda\Desktop\CSE499A_Final_Report_CardioRisk_AI.docx"
DOCX_HRP = r"c:\Users\tonda\Desktop\hrp\Presentations\Reports\CSE499A_Final_Report_CardioRisk_AI.docx"

# Color constants
PRIMARY_HEX = "0D2B45"      # Deep Navy
SECONDARY_HEX = "203C56"    # Slate Navy
ACCENT_HEX = "008080"       # Teal Accent
MUTED_HEX = "5A6B7C"        # Slate Gray
LIGHT_BG_HEX = "F4F6F9"     # Light Gray Background
CALLOUT_BG_HEX = "EBF3FA"   # Soft Blue Callout
BORDER_HEX = "D1D5DB"       # Light Border
SUCCESS_HEX = "1B873F"      # Medical Green
DANGER_HEX = "C53030"       # Crimson Alert

COLOR_PRIMARY = RGBColor(13, 43, 69)
COLOR_SECONDARY = RGBColor(32, 60, 86)
COLOR_ACCENT = RGBColor(0, 128, 128)
COLOR_MUTED = RGBColor(90, 107, 124)
COLOR_TEXT = RGBColor(30, 30, 30)

def set_cell_shading(cell, color_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color_hex=BORDER_HEX, sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color_hex}"/>'
        f'  <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color_hex}"/>'
        f'  <w:left w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'  <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color_hex}"/>'
        f'  <w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def format_cell_text(cell, text, bold=False, italic=False, color=COLOR_TEXT, font_size=9.5, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.1
    run = p.add_run(str(text))
    run.font.name = 'Times New Roman'
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    return run

def add_styled_heading(doc, text, level):
    h = doc.add_heading(text, level=level)
    h.paragraph_format.keep_with_next = True
    run = h.runs[0]
    run.font.name = 'Times New Roman'
    if level == 1:
        run.font.size = Pt(17)
        run.font.bold = True
        run.font.color.rgb = COLOR_PRIMARY
        h.paragraph_format.space_before = Pt(18)
        h.paragraph_format.space_after = Pt(6)
    elif level == 2:
        run.font.size = Pt(13.5)
        run.font.bold = True
        run.font.color.rgb = COLOR_SECONDARY
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after = Pt(4)
    elif level == 3:
        run.font.size = Pt(11.5)
        run.font.bold = True
        run.font.italic = True
        run.font.color.rgb = COLOR_ACCENT
        h.paragraph_format.space_before = Pt(10)
        h.paragraph_format.space_after = Pt(3)
    return h

def add_p(doc, text, bold_prefix="", space_after=6, line_spacing=1.15):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Times New Roman'
        r_pre.font.size = Pt(11)
        r_pre.font.bold = True
        r_pre.font.color.rgb = COLOR_PRIMARY
    r_body = p.add_run(text)
    r_body.font.name = 'Times New Roman'
    r_body.font.size = Pt(11)
    r_body.font.color.rgb = COLOR_TEXT
    return p

def add_bullet(doc, text, bold_prefix="", level=0):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Times New Roman'
        r_pre.font.size = Pt(10.5)
        r_pre.font.bold = True
        r_pre.font.color.rgb = COLOR_PRIMARY
    r_body = p.add_run(text)
    r_body.font.name = 'Times New Roman'
    r_body.font.size = Pt(10.5)
    r_body.font.color.rgb = COLOR_TEXT
    return p

def add_callout(doc, title, text, bg_hex=CALLOUT_BG_HEX, border_hex="1A568C"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_shading(cell, bg_hex)
    set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
    
    # Left border only
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'  <w:left w:val="single" w:sz="24" w:space="0" w:color="{border_hex}"/>'
        f'  <w:top w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'  <w:bottom w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(tcBorders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    r_title = p.add_run(f"★ {title}\n")
    r_title.font.name = 'Times New Roman'
    r_title.font.size = Pt(10.5)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(26, 86, 140)
    
    r_text = p.add_run(text)
    r_text.font.name = 'Times New Roman'
    r_text.font.size = Pt(10.0)
    r_text.font.color.rgb = COLOR_TEXT
    
    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_before = Pt(0)
    p_spacer.paragraph_format.space_after = Pt(4)

def add_formula_block(doc, label, equation_str, explanation):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_shading(cell, "F9FAFB")
    set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
    
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'  <w:left w:val="single" w:sz="18" w:space="0" w:color="008080"/>'
        f'  <w:top w:val="single" w:sz="4" w:space="0" w:color="E5E7EB"/>'
        f'  <w:right w:val="single" w:sz="4" w:space="0" w:color="E5E7EB"/>'
        f'  <w:bottom w:val="single" w:sz="4" w:space="0" w:color="E5E7EB"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(tcBorders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    r_lbl = p.add_run(f"[{label}]\n")
    r_lbl.font.name = 'Times New Roman'
    r_lbl.font.size = Pt(9.5)
    r_lbl.font.bold = True
    r_lbl.font.color.rgb = COLOR_ACCENT
    
    r_eq = p.add_run(f"{equation_str}\n")
    r_eq.font.name = 'Consolas'
    r_eq.font.size = Pt(10.0)
    r_eq.font.bold = True
    r_eq.font.color.rgb = RGBColor(13, 43, 69)
    
    r_exp = p.add_run(f"Where: {explanation}")
    r_exp.font.name = 'Times New Roman'
    r_exp.font.size = Pt(9.0)
    r_exp.font.italic = True
    r_exp.font.color.rgb = COLOR_MUTED
    
    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_after = Pt(4)

def add_figure(doc, img_filename, caption, width_in=5.8):
    img_path = os.path.join(DEMO_DIR, img_filename)
    if not os.path.exists(img_path):
        print(f"[WARN] Image not found: {img_path}")
        return
    
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(8)
    p_img.paragraph_format.space_after = Pt(2)
    p_img.paragraph_format.keep_with_next = True
    
    run_img = p_img.add_run()
    run_img.add_picture(img_path, width=Inches(width_in))
    
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(2)
    p_cap.paragraph_format.space_after = Pt(10)
    
    r_cap = p_cap.add_run(caption)
    r_cap.font.name = 'Times New Roman'
    r_cap.font.size = Pt(9.5)
    r_cap.font.italic = True
    r_cap.font.color.rgb = COLOR_MUTED

def add_code_listing(doc, listing_label, caption, code_lines_text, explanation_text=""):
    p_cap = doc.add_paragraph()
    p_cap.paragraph_format.space_before = Pt(10)
    p_cap.paragraph_format.space_after = Pt(3)
    p_cap.paragraph_format.keep_with_next = True
    
    r_lbl = p_cap.add_run(f"{listing_label}: ")
    r_lbl.font.name = 'Times New Roman'
    r_lbl.font.size = Pt(10.0)
    r_lbl.font.bold = True
    r_lbl.font.color.rgb = COLOR_PRIMARY
    
    r_cap = p_cap.add_run(caption)
    r_cap.font.name = 'Times New Roman'
    r_cap.font.size = Pt(9.5)
    r_cap.font.italic = True
    r_cap.font.color.rgb = COLOR_TEXT
    
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_shading(cell, "F8FAFC")
    set_cell_margins(cell, top=90, bottom=90, left=130, right=130)
    
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'  <w:left w:val="single" w:sz="20" w:space="0" w:color="008080"/>'
        f'  <w:top w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>'
        f'  <w:right w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>'
        f'  <w:bottom w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(tcBorders)
    
    p_code = cell.paragraphs[0]
    p_code.paragraph_format.space_before = Pt(0)
    p_code.paragraph_format.space_after = Pt(0)
    p_code.paragraph_format.line_spacing = 1.05
    
    r_code = p_code.add_run(code_lines_text.strip())
    r_code.font.name = 'Consolas'
    r_code.font.size = Pt(8.5)
    r_code.font.color.rgb = RGBColor(15, 23, 42)
    
    if explanation_text:
        p_exp = doc.add_paragraph()
        p_exp.paragraph_format.space_before = Pt(4)
        p_exp.paragraph_format.space_after = Pt(8)
        p_exp.paragraph_format.line_spacing = 1.15
        
        r_ex_lbl = p_exp.add_run("Algorithmic Rationale & Clinical Role: ")
        r_ex_lbl.font.name = 'Times New Roman'
        r_ex_lbl.font.size = Pt(9.5)
        r_ex_lbl.font.bold = True
        r_ex_lbl.font.color.rgb = COLOR_SECONDARY
        
        r_ex = p_exp.add_run(explanation_text)
        r_ex.font.name = 'Times New Roman'
        r_ex.font.size = Pt(9.5)
        r_ex.font.color.rgb = COLOR_TEXT

def add_toc_item(doc, title, page_num, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5 if level == 1 else 1.5)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    
    if level == 1:
        p.paragraph_format.left_indent = Inches(0.0)
        font_size = 10.5
        bold = True
        color = COLOR_PRIMARY
    elif level == 2:
        p.paragraph_format.left_indent = Inches(0.28)
        font_size = 9.5
        bold = False
        color = COLOR_SECONDARY
    else: # level 3
        p.paragraph_format.left_indent = Inches(0.56)
        font_size = 9.0
        bold = False
        color = COLOR_TEXT
        
    pPr = p._p.get_or_add_pPr()
    tabs = parse_xml(
        r'<w:tabs xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        r'<w:tab w:val="right" w:leader="dot" w:pos="9360"/>'
        r'</w:tabs>'
    )
    pPr.append(tabs)
    
    r_title = p.add_run(title)
    r_title.font.name = 'Times New Roman'
    r_title.font.size = Pt(font_size)
    r_title.font.bold = bold
    r_title.font.color.rgb = color
    
    r_tab = p.add_run(f"\t{page_num}")
    r_tab.font.name = 'Times New Roman'
    r_tab.font.size = Pt(font_size)
    r_tab.font.bold = bold
    r_tab.font.color.rgb = color
    return p

print("Helper definitions loaded successfully.")

